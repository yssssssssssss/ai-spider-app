import asyncio
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import httpx
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

SPEC = {
    "title": "双入口排列", "version": "1", "source": "测试规范.md",
    "document": "当存在两个入口时，两个入口应上下排列；其他数量不适用。",
    "rules": [{"id": "3.7", "clause": "第 3.7 节", "expected": "两个入口应上下排列；其他数量不适用"}],
}


def client_patch(handler):
    original = httpx.AsyncClient
    return patch("app.services.llm_analyzer.httpx.AsyncClient", side_effect=lambda: original(transport=httpx.MockTransport(handler)))


def model_reply(content, model="effective-model"):
    return httpx.Response(200, json={"model": model, "choices": [{"message": {"content": json.dumps(content) if isinstance(content, dict) else content}}]})


def make_analyzer():
    from app.services.llm_analyzer import LLMAnalyzer
    analyzer = LLMAnalyzer()
    analyzer.providers = [{"name": "test", "base_url": "https://test.invalid/v1", "model": "configured-model", "api_key": "secret-for-test"}]
    return analyzer


class InferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = str(Path(self.temp.name) / "capture.png")
        PILImage.new("RGB", (200, 400), "white").save(self.path)

    def tearDown(self):
        self.temp.cleanup()

    def test_target_verdict_requires_boolean_json_and_preserves_false(self):
        analyzer = make_analyzer()
        for content in ("不符合目标页", "符合目标页", '{"is_target":"false"}', '{"is_target":null}', '{"is_target":0}', "[]"):
            with self.subTest(content=content), client_patch(lambda request: model_reply(content)):
                with self.assertRaises(ValueError):
                    asyncio.run(analyzer.is_target_page(self.path))
        with client_patch(lambda request: model_reply({"is_target": False, "reason": "登录页"})):
            self.assertEqual(asyncio.run(analyzer.is_target_page(self.path)), (False, "登录页"))

    def test_fallback_records_failed_provider_and_effective_response_model_without_secrets(self):
        from app.services.model_trace import capture_model_calls
        analyzer = make_analyzer()
        analyzer.providers.insert(0, {**analyzer.providers[0], "name": "unavailable", "base_url": "https://bad.invalid/v1"})
        def handler(request):
            return httpx.Response(400, json={"error": "no provider"}) if request.url.host == "bad.invalid" else model_reply({"is_target": True})
        with capture_model_calls("target_validation") as calls, client_patch(handler):
            self.assertTrue(asyncio.run(analyzer.is_target_page(self.path))[0])
        self.assertEqual([call["status"] for call in calls], ["failed", "success"])
        self.assertEqual(calls[-1]["response_model"], "effective-model")
        self.assertEqual(calls[-1]["requested_model"], "configured-model")
        self.assertNotIn("secret-for-test", json.dumps(calls))
        self.assertTrue(all(call["duration_ms"] >= 0 for call in calls))

    def test_concurrent_requests_do_not_share_trace_records(self):
        from app.services.model_trace import capture_model_calls
        async def run(name):
            with capture_model_calls(name) as calls:
                await make_analyzer().is_target_page(self.path)
            return calls
        async def both():
            return await asyncio.gather(run("scene-a"), run("scene-b"))
        with client_patch(lambda request: model_reply({"is_target": True})):
            calls = asyncio.run(both())
        self.assertEqual([[call["purpose"] for call in group] for group in calls], [["scene-a"], ["scene-b"]])

    def test_custom_output_schema_is_enforced(self):
        from jsonschema.exceptions import ValidationError
        skill = SimpleNamespace(skill_type="analysis", prompt="test", output_schema_json={"type": "object", "required": ["valid"], "properties": {"valid": {"type": "boolean"}}})
        with client_patch(lambda request: model_reply({"valid": "false"})):
            with self.assertRaises(ValidationError):
                asyncio.run(make_analyzer().analyze_with_skill(self.path, skill))

    def test_output_schema_cannot_fetch_remote_refs(self):
        from app.services.inspection import validate_output_schema
        with self.assertRaises(ValueError):
            validate_output_schema({"$ref": "http://127.0.0.1:8000/private"})

    def test_excel_preserves_long_specifications_and_history(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from app.services.exporter import excel_bytes
        long_spec = {**SPEC, "document": "规范🙂" * 20000}
        inspection = {"specification": long_spec, "checks": [{"rule_id": "3.7", "status": "fail"}]}
        payload = {"task": {"name": "long-export"}, "runs": [{"analysis_skills_snapshot_json": [long_spec]}],
                   "images": [{"id": "image-1", "file_path": "image.png", "analyses": [{"inspection_json": inspection}],
                               "analysis_history": [{"analysis": {"inspection_json": inspection}}]}]}
        workbook = load_workbook(BytesIO(excel_bytes(payload)))
        for sheet_name, field, expected in (("runs", "analysis_skills_snapshot_json", [long_spec]),
                                            ("analysis", "inspection_json", inspection),
                                            ("analysis_history", "analysis", {"inspection_json": inspection})):
            sheet = workbook[sheet_name]
            columns = [i for i, cell in enumerate(sheet[1]) if str(cell.value).startswith(field + "__part_")]
            value = "".join(sheet[2][i].value for i in columns)
            self.assertEqual(json.loads(value), expected)

    def test_excel_text_chunks_are_not_executable_formulas(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from app.services.exporter import excel_bytes
        value = "a" * 16000 + "=1+1"
        payload = {"task": {"name": value}, "runs": [], "images": []}
        workbook = load_workbook(BytesIO(excel_bytes(payload)), data_only=True)
        values = [cell.value for cell in workbook["overview"][2]]
        self.assertIn("=1+1", values)
        workbook = load_workbook(BytesIO(excel_bytes(payload)))
        self.assertTrue(all(cell.data_type != "f" for cell in workbook["overview"][2]))

    def test_missing_and_invalid_evidence_are_uncertain(self):
        from app.services.inspection import Specification, build_inspection
        image = SimpleNamespace(id=uuid4(), source_app="京东", scenario="促销页面", task_run_id=None, device_id=None)
        for result in ({"checks": []}, {"checks": [{"rule_id": "3.7", "status": "pass", "observed": "两个入口", "reason": "上下排列", "bbox_norm": [0, 0, 1001, 500]}]}):
            report = build_inspection(result, Specification.model_validate(SPEC), image, [], {}, (200, 400))
            self.assertEqual(report["summary"]["uncertain"], 1)
            self.assertEqual(report["checks"][0]["evidence"], [])

    def test_streaming_phone_calls_keep_effective_model_and_close_stream(self):
        from worker.model_trace import trace_phone_client, read_model_calls
        class Stream:
            closed = False
            def __iter__(self):
                yield SimpleNamespace(model="autoglm-effective", choices=[])
            def close(self):
                self.closed = True
        stream = Stream()
        client = SimpleNamespace(client=SimpleNamespace(base_url="https://user:password@test.invalid/v1?token=secret", chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kwargs: stream))))
        trace_phone_client(client, self.temp.name)
        list(client.client.chat.completions.create(model="AutoGLM-Phone-9B", stream=True))
        calls = read_model_calls(self.temp.name)
        self.assertTrue(stream.closed)
        self.assertEqual(calls[0]["response_model"], "autoglm-effective")
        self.assertEqual(calls[0]["endpoint_host"], "test.invalid")
        self.assertNotIn("password", json.dumps(calls))

    def test_trace_io_failure_preserves_original_model_error(self):
        from worker.model_trace import trace_phone_client
        def fail(**kwargs):
            raise RuntimeError("model is unavailable")
        client = SimpleNamespace(client=SimpleNamespace(base_url="https://test.invalid/v1", chat=SimpleNamespace(completions=SimpleNamespace(create=fail))))
        trace_phone_client(client, self.temp.name)
        with patch("worker.model_trace.Path.open", side_effect=OSError("disk unavailable")):
            with self.assertRaisesRegex(RuntimeError, "model is unavailable"):
                client.client.chat.completions.create(model="configured", stream=True)

    def test_truncated_trace_is_explicit_without_breaking_run_finalization(self):
        from worker.model_trace import read_model_calls
        from app.services.model_trace import read_phone_model_calls
        Path(self.temp.name, "model_calls.jsonl").write_text('{"requested_model":"configured"}\n{"requested_model":', encoding="utf-8")
        for reader in (read_model_calls, read_phone_model_calls):
            calls = reader(self.temp.name)
            self.assertEqual(calls[0]["requested_model"], "configured")
            self.assertEqual(calls[1]["status"], "trace_error")


class InspectionWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app.database import engine, ensure_schema
        if "test" not in engine.url.database and not engine.url.database.endswith("_round1"):
            raise unittest.SkipTest("Integration tests require a disposable test database")
        ensure_schema()

    def setUp(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from app.database import SessionLocal
        from app.routers import analysis_skills, images
        from app.services.auth import get_current_user
        self.db = SessionLocal()
        (ROOT / "data").mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / "data")
        self.user = SimpleNamespace(id=uuid4(), role="operator")
        app = FastAPI()
        app.include_router(analysis_skills.router, prefix="/api")
        app.include_router(images.router, prefix="/api")
        app.dependency_overrides[get_current_user] = lambda: self.user
        self.client = TestClient(app)
        self.tasks, self.skill_ids = [], []

    def tearDown(self):
        from app import models
        for task in self.tasks:
            for image in self.db.query(models.Image).filter(models.Image.task_id == task).all():
                self.db.delete(image)
            self.db.delete(self.db.get(models.Task, task))
        self.db.flush()
        for skill in self.skill_ids:
            self.db.delete(self.db.get(models.AnalysisSkill, skill))
        self.db.commit()
        self.db.close()
        self.client.close()
        self.temp.cleanup()

    def create_skill(self):
        response = self.client.post("/api/admin/analysis-skills", json={"name": "可复用规范", "prompt": "检查规范", "skill_type": "inspection", "specification_json": SPEC})
        self.assertEqual(response.status_code, 200, response.text)
        self.skill_ids.append(response.json()["id"])
        return response.json()

    def test_one_spec_reused_across_scenes_and_history_survives_rule_edits_and_failure(self):
        from app import crud, models, schemas
        from app.routers import images
        from app.services.exporter import task_export_payload, excel_bytes
        from io import BytesIO
        from openpyxl import load_workbook
        skill = self.create_skill()
        image_ids = []
        observed = {"checks": [{"rule_id": "3.7", "status": "pass", "observed": "两个入口上下排列", "reason": "证据显示上下排列", "bbox_norm": [0, 0, 500, 500]}]}
        def handler(request):
            payload = json.loads(request.content)
            prompt = payload["messages"][0]["content"][0]["text"]
            return model_reply({"is_target": True, "reason": "已到达"}) if '"is_target"' in prompt else model_reply(observed)
        for index, scene in enumerate(("新品页", "促销页面")):
            task = crud.create_task(self.db, name="inspection-test", keyword="", target_app="京东", target_scenario=scene)
            self.tasks.append(task.id)
            task.analysis_skill_ids = [skill["id"]]
            self.db.commit()
            path = Path(self.temp.name) / f"scene-{index}.png"
            PILImage.new("RGB", (200, 400), "white").save(path)
            image = crud.create_image(self.db, schemas.ImageCreate(file_path=str(path.relative_to(ROOT)), task_id=task.id, source_app="京东", scenario=scene))
            image_ids.append(image.id)
            with patch.object(images, "analyzer", make_analyzer()), patch.object(images, "_embed_skill_result", new=AsyncMock()), client_patch(handler):
                asyncio.run(images._analyze_and_embed(image.id))
            self.db.expire_all()
            analysis = crud.get_analysis_by_image_and_skill(self.db, image.id, skill["id"])
            self.assertEqual(analysis.inspection_json["checks"][0]["status"], "pass")
            self.assertEqual(analysis.inspection_json["scene"]["name"], scene)
            self.assertEqual(analysis.provenance_json["model_calls"][-1]["response_model"], "effective-model")
        edited = {**SPEC, "version": "2", "document": "两个入口应左右排列", "rules": [{"id": "3.7", "clause": "第 3.7 节", "expected": "两个入口应左右排列"}]}
        response = self.client.patch(f"/api/admin/analysis-skills/{skill['id']}", json={"specification_json": edited})
        self.assertEqual(response.json()["version"], 2)
        # A no-op edit must not create a new version.
        response = self.client.patch(f"/api/admin/analysis-skills/{skill['id']}", json={"specification_json": edited})
        self.assertEqual(response.json()["version"], 2)
        # Reanalysis fails: stale success/raw JSON must be cleared, previous report retained.
        with patch.object(images, "analyzer", make_analyzer()), client_patch(lambda req: model_reply("不符合")):
            asyncio.run(images._analyze_and_embed(image_ids[0]))
        history = self.client.get(f"/api/images/{image_ids[0]}/analysis-history").json()
        self.db.expire_all()
        analysis = crud.get_analysis_by_image_and_skill(self.db, image_ids[0], skill["id"])
        self.assertEqual(analysis.status, "failed")
        self.assertIsNone(analysis.result_json)
        self.assertEqual(analysis.provenance_json["skill"]["version"], 2)
        self.assertEqual(history[0]["analysis"]["inspection_json"]["specification"]["version"], "1")
        payload = task_export_payload(self.db, self.tasks[0])
        self.assertIsNone(payload["images"][0]["analyses"][0]["inspection_json"])
        self.assertTrue(payload["images"][0]["analysis_history"])
        workbook = load_workbook(BytesIO(excel_bytes(payload)))
        self.assertGreater(workbook["analysis_history"].max_row, 1)
        self.assertIn("provenance_json", [cell.value for cell in workbook["analysis"][1]])
        self.user.role = "viewer"
        self.assertEqual(self.client.get(f"/api/images/{image_ids[0]}/analysis-history").status_code, 404)

    def test_invalid_spec_and_remote_schema_rejected_at_api_boundary(self):
        for body in (
            {"skill_type": "inspection", "specification_json": {**SPEC, "rules": SPEC["rules"] * 2}},
            {"output_schema_json": {"$ref": "https://example.invalid/private"}},
        ):
            response = self.client.post("/api/admin/analysis-skills", json={"name": "invalid", "prompt": "test", **body})
            self.assertEqual(response.status_code, 400, response.text)

    def test_run_uses_frozen_rules_after_skill_and_task_bindings_change(self):
        from app import crud, models, schemas
        from app.routers import images
        skill = self.create_skill()
        task = crud.create_task(self.db, name="snapshot-test", keyword="", target_app="京东", target_scenario="新品页")
        self.tasks.append(task.id)
        task.analysis_skill_ids = [skill["id"]]
        self.db.commit()
        run = crud.create_task_run(self.db, task.id)
        changed = {**SPEC, "version": "2", "document": "新的规范文本"}
        self.client.patch(f"/api/admin/analysis-skills/{skill['id']}", json={"specification_json": changed})
        task.analysis_skill_ids = []
        self.db.commit()
        path = Path(self.temp.name) / "run.png"
        PILImage.new("RGB", (1080, 2400), "white").save(path)
        image = crud.create_image(self.db, schemas.ImageCreate(file_path=str(path.relative_to(ROOT)), task_id=task.id, task_run_id=run.id, source_app="京东", scenario="新品页"))
        observed = {"checks": [{"rule_id": "3.7", "status": "fail", "observed": "左右排列", "reason": "应上下排列", "bbox_norm": [500, 100, 900, 500]}]}
        def handler(request):
            prompt = json.loads(request.content)["messages"][0]["content"][0]["text"]
            return model_reply({"is_target": True}) if '"is_target"' in prompt else model_reply(observed)
        with patch.object(images, "analyzer", make_analyzer()), patch.object(images, "_embed_skill_result", new=AsyncMock()), client_patch(handler):
            asyncio.run(images._analyze_and_embed(image.id))
        self.db.expire_all()
        analysis = crud.get_analysis_by_image_and_skill(self.db, image.id, skill["id"])
        self.assertEqual(analysis.inspection_json["specification"]["version"], "1")
        self.assertEqual(analysis.provenance_json["binding_mode"], "run_snapshot")
        self.assertEqual(analysis.inspection_json["summary"]["fail"], 1)
        self.assertEqual(analysis.inspection_json["checks"][0]["evidence"][0]["bbox_px"], [540, 240, 972, 1200])

    def test_system_skill_cannot_be_converted_to_inspection(self):
        from app import models
        system = self.db.query(models.AnalysisSkill).filter(models.AnalysisSkill.is_system.is_(True)).first()
        response = self.client.patch(f"/api/admin/analysis-skills/{system.id}", json={"skill_type": "inspection", "specification_json": SPEC})
        self.assertEqual(response.status_code, 400)

    def test_missing_skill_cannot_silently_create_default_analysis_run(self):
        from app import crud, models
        task = crud.create_task(self.db, name="missing-skill-test", keyword="", target_app="京东", target_scenario="新品页")
        self.tasks.append(task.id)
        task.analysis_skill_ids = [uuid4()]
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "Skill 不存在"):
            crud.create_task_run(self.db, task.id)
        self.assertEqual(self.db.query(models.TaskRun).filter(models.TaskRun.task_id == task.id).count(), 0)

    def test_system_output_schema_is_applied_to_actual_result(self):
        from app import crud, models, schemas
        from app.routers import images
        system = self.db.query(models.AnalysisSkill).filter(models.AnalysisSkill.is_system.is_(True)).first()
        original_schema, original_version = system.output_schema_json, system.version
        try:
            self.client.patch(f"/api/admin/analysis-skills/{system.id}", json={"output_schema_json": {"required": ["checked"]}})
            task = crud.create_task(self.db, name="system-schema-test", keyword="", target_app="京东", target_scenario="新品页")
            self.tasks.append(task.id)
            task.analysis_skill_ids = [system.id]
            self.db.commit()
            path = Path(self.temp.name) / "system-schema.png"
            PILImage.new("RGB", (200, 400), "white").save(path)
            image = crud.create_image(self.db, schemas.ImageCreate(file_path=str(path.relative_to(ROOT)), task_id=task.id))
            def handler(request):
                prompt = json.loads(request.content)["messages"][0]["content"][0]["text"]
                return model_reply({"is_target": True}) if '"is_target"' in prompt else model_reply({"design_analysis": "design", "ops_analysis": "ops"})
            with patch.object(images, "analyzer", make_analyzer()), client_patch(handler):
                asyncio.run(images._analyze_and_embed(image.id))
            self.db.expire_all()
            analysis = crud.get_analysis_by_image_and_skill(self.db, image.id, system.id)
            self.assertEqual(analysis.status, "failed")
            self.assertIsNone(analysis.result_json)
        finally:
            self.db.refresh(system)
            system.output_schema_json, system.version = original_schema, original_version
            self.db.commit()
